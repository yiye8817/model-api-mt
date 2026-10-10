"""Local document decoding and deterministic text chunking."""
from dataclasses import dataclass
import io
from pathlib import Path


class DocumentError(ValueError):
    """A document cannot be parsed or chunked for a user-facing reason."""


@dataclass(frozen=True)
class ParsedDocument:
    text: str
    filename: str
    media_type: str


def _decode_text(content: bytes) -> str:
    if not isinstance(content, (bytes, bytearray)):
        raise DocumentError('文档内容必须是二进制数据')
    try:
        value = bytes(content).decode('utf-8-sig')
    except UnicodeDecodeError as exc:
        raise DocumentError('文本文档必须使用 UTF-8 编码') from exc
    value = value.strip()
    if not value:
        raise DocumentError('文档内容不能为空')
    return value


def _parse_pdf(content: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise DocumentError('解析 PDF 需要安装 pypdf') from exc
    try:
        reader = PdfReader(io.BytesIO(content))
        value = '\n\n'.join((page.extract_text() or '') for page in reader.pages).strip()
    except Exception as exc:
        raise DocumentError(f'PDF 解析失败：{exc}') from exc
    if not value:
        raise DocumentError('PDF 未提取到文本内容')
    return value


def _parse_docx(content: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise DocumentError('解析 DOCX 需要安装 python-docx') from exc
    try:
        document = Document(io.BytesIO(content))
        value = '\n'.join(paragraph.text for paragraph in document.paragraphs).strip()
    except Exception as exc:
        raise DocumentError(f'DOCX 解析失败：{exc}') from exc
    if not value:
        raise DocumentError('DOCX 未提取到文本内容')
    return value


def parse_document(content: bytes, filename: str, fmt: str) -> ParsedDocument:
    """Decode a local document into plain text.

    Structured JSON imports are intentionally handled by datasets.convert; this
    function is for file-like documents that need textual indexing.
    """
    if not isinstance(filename, str) or not filename.strip():
        raise DocumentError('文件名不能为空')
    fmt = str(fmt or '').strip().lower()
    name = Path(filename).name or filename
    media_type = 'text/plain'
    if fmt in ('text', 'txt'):
        value = _decode_text(content)
    elif fmt in ('markdown', 'md'):
        value = _decode_text(content)
        media_type = 'text/markdown'
    elif fmt == 'pdf' or name.lower().endswith('.pdf'):
        value = _parse_pdf(content)
        media_type = 'application/pdf'
    elif fmt == 'docx' or name.lower().endswith('.docx'):
        value = _parse_docx(content)
        media_type = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    else:
        raise DocumentError(f'不支持的文档格式：{fmt or name}')
    return ParsedDocument(value, name, media_type)


def chunk_text(text: str, size: int = 800, overlap: int = 120) -> list[str]:
    """Split text into stable, non-empty character windows with overlap."""
    if not isinstance(text, str) or not text.strip():
        raise DocumentError('文档内容不能为空')
    if type(size) is not int or not 200 <= size <= 4000:
        raise DocumentError('chunk_size 必须是 200–4000 的整数')
    if type(overlap) is not int or overlap < 0 or overlap >= size:
        raise DocumentError('chunk_overlap 必须是小于 chunk_size 的非负整数')
    value = text.strip()
    step = size - overlap
    chunks = []
    for start in range(0, len(value), step):
        chunk = value[start:start + size].strip()
        if chunk:
            chunks.append(chunk)
        if start + size >= len(value):
            break
    return chunks
