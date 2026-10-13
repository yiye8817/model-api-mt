import io
import os
import time

import pytest


def test_chunk_limits_and_original_positions(tmp_path):
    from local_knowledge.parsers import parse_file
    from local_knowledge.text import chunk_document
    text = '# 标题\n\n' + ('产品 A 保修三年。\n' * 200) + '\n```python\n' + ('print(1)\n' * 250) + '```\n'
    path = tmp_path / 'manual.md'
    path.write_text(text)
    doc = parse_file(path)
    chunks = chunk_document(doc)
    assert len(chunks) > 2
    assert all(len(c.text) <= 1600 for c in chunks)
    assert all(text[c.position.start_char:c.position.end_char] == c.text for c in chunks)
    assert all(c.position.start_line >= 1 for c in chunks)
    assert any(c.continuation_group for c in chunks)


def test_docx_table_positions(tmp_path):
    from docx import Document
    from local_knowledge.parsers import parse_file
    doc = Document()
    doc.add_paragraph('保修说明')
    table = doc.add_table(rows=2,cols=2)
    table.cell(0,0).text='型号'; table.cell(0,1).text='保修期'
    table.cell(1,0).text='A'; table.cell(1,1).text='三年'
    path = tmp_path / 'manual.docx'
    doc.save(path)
    units = parse_file(path).units
    assert units[0].position.paragraph == 1
    assert any(u.position.table == 1 and '三年' in u.text and '保修期' in u.text for u in units)


@pytest.mark.parametrize('kind', ['scanned','encrypted','corrupt'])
def test_scanned_encrypted_corrupt_pdf_reported(tmp_path, kind):
    from pypdf import PdfWriter
    from local_knowledge.parsers import parse_file, ParseError
    path = tmp_path / 'document.pdf'
    if kind == 'corrupt':
        path.write_bytes(b'%PDF broken')
    else:
        w=PdfWriter(); w.add_blank_page(width=100,height=100)
        if kind == 'encrypted': w.encrypt('secret')
        w.write(path)
    with pytest.raises(ParseError):
        parse_file(path)


def test_twenty_mib_and_two_thousand_files_boundaries(tmp_path):
    from local_knowledge.ingest import scan_root
    from local_knowledge.parsers import parse_file, ParseError
    from local_knowledge.contracts import ParseLimits
    path=tmp_path/'large.txt'
    path.write_bytes(b'x' * (20*1024*1024+1))
    with pytest.raises(ParseError, match='file_too_large'):
        parse_file(path)
    path.write_bytes(b'abc')
    assert parse_file(path, ParseLimits(max_file_bytes=3)).units[0].text == 'abc'
    with pytest.raises(ParseError): parse_file(path, ParseLimits(max_file_bytes=2))
    path.unlink()
    for i in range(2001): (tmp_path / f'{i:04}.txt').write_text('sample')
    result=scan_root(tmp_path)
    assert len(result.files) == 2000
    assert any(item['reason']=='file_count_limit' for item in result.failed)


def test_symlink_and_sensitive_names_rejected(tmp_path):
    from local_knowledge.ingest import scan_root
    root=tmp_path/'selected'; root.mkdir()
    outside=tmp_path/'outside.txt'; outside.write_text('outside')
    (root/'escape.txt').symlink_to(outside)
    (root/'.env').write_text('TOKEN=secret')
    (root/'normal.md').write_text('allowed')
    (root/'.notes.md').write_text('hidden')
    result=scan_root(root)
    assert [p.name for p in result.files] == ['normal.md']
    assert 'outside_root' in [i['reason'] for i in result.skipped]
    assert [p.name for p in scan_root(root,[root/'.notes.md']).files] == ['.notes.md']


@pytest.mark.parametrize('name',['../escape.txt','/absolute.txt','..\\escape.txt','C:\\escape.txt'])
def test_upload_traversal_rejected(tmp_path,name):
    from local_knowledge.ingest import ImportManager
    from local_knowledge.store import Store
    from local_knowledge.contracts import UploadItem
    manager=ImportManager(Store(tmp_path/'data'))
    try:
        job=manager.enqueue({'kind':'upload','purpose':'knowledge'})
        manager.append_upload(job,[UploadItem(name,io.BytesIO(b'x'))],True)
        manager.jobs.queue.join()
        assert manager.snapshot(job)['failed']==1
        assert not (tmp_path/'escape.txt').exists()
    finally: manager.close()


def test_parser_timeout_kills_worker(tmp_path):
    import multiprocessing
    from local_knowledge.parsers import parse_file, ParseError
    from local_knowledge.contracts import ParseLimits
    path=tmp_path/'text.txt'; path.write_text('text')
    before={p.pid for p in multiprocessing.active_children()}
    with pytest.raises(ParseError,match='parse_timeout'):
        parse_file(path,ParseLimits(timeout_seconds=.001))
    assert {p.pid for p in multiprocessing.active_children()} == before


def test_file_changed_with_same_size_and_mtime(tmp_path):
    from local_knowledge.ingest import file_source, check_file_freshness
    path=tmp_path/'manual.txt'; path.write_text('first')
    source=file_source(path,tmp_path)
    stat=path.stat()
    path.write_text('other'); os.utime(path,ns=(stat.st_atime_ns,stat.st_mtime_ns))
    assert not check_file_freshness(source)


def test_upload_records_bad_items_and_imports_good_ones(tmp_path):
    from local_knowledge.ingest import ImportManager
    from local_knowledge.store import Store
    from local_knowledge.contracts import UploadItem
    from tests.knowledge.test_jobs import wait_job
    store=Store(tmp_path/'data');manager=ImportManager(store,lambda sid,rev,b:store.activate_revision(sid,rev))
    try:
        job_id=manager.enqueue({'kind':'upload','purpose':'knowledge'})
        manager.append_upload(job_id,[UploadItem('good.txt',io.BytesIO(b'allowed')),UploadItem('../bad.txt',io.BytesIO(b'bad')),UploadItem('.env',io.BytesIO(b'secret'))],True,'batch')
        job=wait_job(manager,job_id)
        assert job['succeeded']==1 and job['failed']+job['skipped']==2
        assert len(store.visible_chunks())==1
    finally:manager.close()


def test_retried_upload_batch_is_idempotent_with_rejected_items(tmp_path):
    from local_knowledge.ingest import ImportManager
    from local_knowledge.store import Store
    from local_knowledge.contracts import UploadItem
    manager=ImportManager(Store(tmp_path/'data'))
    try:
        job=manager.enqueue({'kind':'upload'})
        def items():return [UploadItem('good.txt',io.BytesIO(b'good')),UploadItem('../bad.txt',io.BytesIO(b'bad'))]
        manager.append_upload(job,items(),False,'batch')
        manager.append_upload(job,items(),False,'batch')
        assert manager.snapshot(job)['total']==2 and len(manager.snapshot(job)['payload']['files'])==1
    finally:manager.close()
