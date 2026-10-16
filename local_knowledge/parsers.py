"""Document parsing runs in an independent, bounded worker."""
import codecs
import multiprocessing
import re
import time
import zipfile
from collections import Counter
from dataclasses import asdict
from pathlib import Path

from .contracts import ParsedDocument, ParseLimits, Position, TextUnit
from .text import text_units

TEXT_SUFFIXES=set('.txt .md .markdown .py .js .jsx .ts .tsx .c .h .cpp .hpp .java .go .rs .sh .html .css .sql .json .csv .yaml .yml .toml'.split())


class ParseError(RuntimeError):
    pass


def _pdf_images(page,limits):
    """Measure image coverage through graphics transforms and nested Form resources."""
    from pypdf.generic import ContentStream
    threshold=float(page.mediabox.width*page.mediabox.height)*.5
    budget=0;has_images=False
    def determinant(matrix):return abs(float(matrix[0])*float(matrix[3])-float(matrix[1])*float(matrix[2]))
    def walk(stream,resources,scale,ancestors):
        nonlocal budget,has_images
        if stream is None:return False
        budget+=len(stream.get_data())
        if budget>limits.max_pdf_page_bytes:raise ParseError('pdf_page_expanded_limit')
        if len(ancestors)>16:raise ParseError('pdf_nested_content_limit')
        if hasattr(resources,'get_object'):resources=resources.get_object()
        objects=resources.get('/XObject',{})
        if hasattr(objects,'get_object'):objects=objects.get_object()
        stack=[]
        for args,operator in ContentStream(stream,page.pdf).operations:
            if operator==b'q':stack.append(scale)
            elif operator==b'Q':scale=stack.pop() if stack else scale
            elif operator==b'cm':scale*=determinant(args)
            elif operator==b'INLINE IMAGE':
                has_images=True
                if scale>=threshold:return True
            elif operator==b'Do' and args:
                obj=objects.get(args[0])
                if obj is None:continue
                obj=obj.get_object()
                if obj.get('/Subtype')=='/Image':
                    has_images=True
                    if scale>=threshold:return True
                if obj.get('/Subtype')=='/Form' and id(obj) not in ancestors:
                    factor=determinant(obj.get('/Matrix',[1,0,0,1,0,0]))
                    if walk(obj,obj.get('/Resources',resources),scale*factor,ancestors|{id(obj)}):return True
        return False
    large=walk(page.get_contents(),page.get('/Resources',{}),1,set())
    return large,has_images


def _parse_worker(sender,path,limits,encoding,allow_scan_placeholders=False):
    try:
        try:
            import resource
            resource.setrlimit(resource.RLIMIT_AS,(limits.max_memory_bytes,limits.max_memory_bytes))
        except ImportError:
            raise ParseError('memory_limit_unavailable')
        path=Path(path)
        if path.suffix.lower()=='.pdf':
            from pypdf import PdfReader
            from pypdf.errors import PdfReadError
            reader=PdfReader(path)
            if reader.is_encrypted: raise ParseError('encrypted_pdf')
            units=[]; offset=0; warnings=[]
            for page_no,page in enumerate(reader.pages,1):
                contents=page.get_contents()
                if contents is not None and len(contents.get_data())>limits.max_pdf_page_bytes:
                    raise ParseError('pdf_page_expanded_limit')
                # Layout extraction does not support every native PDF text form.
                # Keep the ordinary extraction whenever layout fails or loses text.
                text=page.extract_text() or ''
                if offset+len(text)>limits.max_text_chars: raise ParseError('expanded_text_limit')
                if contents is not None:
                    try:
                        layout=page.extract_text(extraction_mode='layout',layout_mode_strip_rotated=False,layout_mode_space_vertically=False) or ''
                    except (KeyError,ValueError,TypeError,NotImplementedError,PdfReadError):
                        layout=''
                    if not (Counter(c for c in text if not c.isspace())-Counter(c for c in layout if not c.isspace())):
                        text=layout
                if sum(bool(re.search(r'\S {3,}\S',line)) for line in text.splitlines()) >= 2:
                    warnings.append(f'第 {page_no} 页含表格或分栏，已保留排版；行列对应关系请对照原页核对。')
                large,images=_pdf_images(page,limits) if allow_scan_placeholders else (False,False)
                scan=allow_scan_placeholders and (large or len(''.join(text.split()))<80 and images)
                units.append(TextUnit(text,Position(kind='pdf_page',page=page_no,start_char=offset,end_char=offset+len(text)), 'scan' if scan else 'paragraph'))
                offset+=len(text)+1
                if offset>limits.max_text_chars: raise ParseError('expanded_text_limit')
            if not any(u.text.strip() or u.kind=='scan' for u in units): raise ParseError('scanned_or_empty_pdf')
            result=ParsedDocument(units,warnings=warnings)
        elif path.suffix.lower()=='.docx':
            with zipfile.ZipFile(path) as archive:
                if sum(i.file_size for i in archive.infolist())>limits.max_pdf_page_bytes:
                    raise ParseError('docx_expanded_limit')
            from docx import Document
            from docx.text.paragraph import Paragraph
            from docx.table import Table
            doc=Document(path); units=[]; offset=0; paragraph=0; table_no=0
            for element in doc.element.body:
                if element.tag.endswith('}p'):
                    paragraph+=1; value=Paragraph(element,doc).text
                    pos=Position(kind='docx',paragraph=paragraph,start_char=offset,end_char=offset+len(value))
                    units.append(TextUnit(value,pos)); offset+=len(value)+1
                elif element.tag.endswith('}tbl'):
                    table_no+=1; table=Table(element,doc)
                    header=' | '.join(c.text for c in table.rows[0].cells) if table.rows else ''
                    for row_no,row in enumerate(table.rows,1):
                        value=' | '.join(c.text for c in row.cells)
                        if row_no>1: value=header+'\n'+value
                        pos=Position(kind='docx',table=table_no,row=row_no,start_char=offset,end_char=offset+len(value))
                        units.append(TextUnit(value,pos,'table')); offset+=len(value)+1
                    for row_no,row in enumerate(table.rows,1):
                        for col_no,cell in enumerate(row.cells,1):
                            value=cell.text
                            pos=Position(kind='docx',table=table_no,row=row_no,column=col_no,start_char=offset,end_char=offset+len(value))
                            units.append(TextUnit(value,pos,'cell')); offset+=len(value)+1
            result=ParsedDocument(units)
        else:
            raw=path.read_bytes()
            codec=encoding or ('utf-16' if raw.startswith((codecs.BOM_UTF16_LE,codecs.BOM_UTF16_BE)) else 'utf-8-sig')
            try: text=raw.decode(codec)
            except (UnicodeError,LookupError) as exc: raise ParseError('unsupported_encoding') from exc
            if '\0' in text: raise ParseError('binary_content')
            result=ParsedDocument(text_units(text))
        if sum(len(u.text) for u in result.units)>limits.max_text_chars: raise ParseError('expanded_text_limit')
        if not any(u.text.strip() or u.kind=='scan' for u in result.units): raise ParseError('empty_document')
        sender.send((True,asdict(result)))
    except Exception as exc:
        sender.send((False,str(exc)[:300]))
    finally: sender.close()


def parse_file(path,limits=None,encoding=None,cancel_event=None,allow_scan_placeholders=False):
    limits=limits or ParseLimits(); path=Path(path)
    if path.stat().st_size>limits.max_file_bytes: raise ParseError('file_too_large')
    if path.suffix.lower() not in TEXT_SUFFIXES|{'.pdf','.docx'}: raise ParseError('unsupported_extension')
    ctx=multiprocessing.get_context('spawn'); receiver,sender=ctx.Pipe(duplex=False)
    worker=ctx.Process(target=_parse_worker,args=(sender,str(path),limits,encoding,allow_scan_placeholders),daemon=True)
    deadline=time.monotonic()+limits.timeout_seconds
    worker.start(); sender.close()
    try:
        while not receiver.poll(min(.02,max(0,deadline-time.monotonic()))):
            if cancel_event and cancel_event.is_set(): raise ParseError('cancelled')
            if time.monotonic()>=deadline: raise ParseError('parse_timeout')
            if not worker.is_alive(): raise ParseError('parse_worker_exited')
        ok,data=receiver.recv()
        if cancel_event and cancel_event.is_set(): raise ParseError('cancelled')
        if not ok: raise ParseError(data)
        return ParsedDocument([TextUnit(u['text'],Position(**u['position']),u['kind'],u['continuation_group']) for u in data['units']],data['nature'],data['warnings'])
    except EOFError as exc: raise ParseError('parse_worker_exited') from exc
    finally:
        receiver.close(); worker.join(.05)
        if worker.is_alive(): worker.terminate(); worker.join(.5)
        if worker.is_alive(): worker.kill(); worker.join(.5)
        worker.close()
